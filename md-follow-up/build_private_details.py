#!/usr/bin/env python3
"""Build locally held, read-only MD work details. NEVER commit records.json to GitHub.
Public snapshot remains aggregate only. Requires authorized NB with mapped W: drive.
"""
import json, os, pathlib, datetime as dt, re, sys
from openpyxl import load_workbook

ROOT=pathlib.Path(__file__).resolve().parent
PRIVATE=pathlib.Path(os.environ.get("LOCALAPPDATA",str(pathlib.Path.home())))/"UAIOS"/"MD-FollowUp-Private"
PRIVATE.mkdir(parents=True,exist_ok=True)
TODAY=dt.date.today()
def value(x):
    if x is None:return ""
    if isinstance(x,(dt.date,dt.datetime)):return x.strftime("%Y-%m-%d")
    return str(x).strip()
def safe(x,maxlen=600):return value(x)[:maxlen]
def parse_due(raw):
    s=value(raw)
    if not s:return None
    if re.match(r"^\d{4}-\d{2}-\d{2}",s):
        try:return dt.date.fromisoformat(s[:10])
        except ValueError:return None
    m=re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})$",s)
    if not m:return None
    y=int(m.group(3))
    if y>=2500:y-=543
    elif y==69 or y==26:y=2026
    elif y<100:y+=2000
    try:return dt.date(y,int(m.group(2)),int(m.group(1)))
    except ValueError:return None
def due_state(s,st,is_due=True):
    if not is_due:return "ไม่ยืนยันกำหนด"
    d=parse_due(s)
    if not d:return "ไม่ระบุวัน"
    if st in ("สำเร็จ","จบ","เข้า QP","ยกเลิก","เปิด Order ใหม่"):return "ปิด/ส่งต่อตามสถานะ"
    if d<TODAY:return "ผ่านกำหนดในตาราง"
    if d==TODAY:return "ครบกำหนดวันนี้"
    return "ยังไม่ถึงกำหนด"
def create(source,rid,title,status,owner="",due="",due_kind="กำหนดส่ง",notes="",ref="",parent="",progress=None,details=None):
    state=due_state(due,status,due_kind in ("กำหนดส่ง","วันที่โอน"))
    a={"source":source,"id":safe(rid,140),"title":safe(title),"status":safe(status,140) or "ไม่ระบุ",
       "owner":safe(owner,180) or "ไม่ระบุ","due":safe(due,75),"due_kind":due_kind,"timeliness":state,
       "notes":safe(notes,600),"reference":safe(ref,200),"parent":safe(parent,250)}
    if progress is not None:a["percent"]=progress
    if details:a["steps"]=details
    return a
conf=json.loads((ROOT/"refresh.config.json").read_text(encoding="utf-8-sig"))
groups={}
warnings=[]
# CB worksheet drilldown. Template rows are not verified project execution.
def cb_code(x):
    m=re.match(r"^\s*(\d{1,2}\.\d{1,3})(?=\s|$)",value(x).replace("\n"," "))
    return m.group(1) if m else ""
def cb_sheet_steps(sheet,template=False):
    rows=[]
    for rn,raw in enumerate(sheet.iter_rows(min_row=5,max_col=10,values_only=True),5):
        cells=list(raw)+[None]*max(0,10-len(raw))
        code=cb_code(cells[1])
        if not code or not value(cells[4]):continue
        check=value(cells[2]).casefold()
        checked=check in ("1","true","yes","done","complete","✓","✔","ผ่าน","เสร็จ")
        status="แม่แบบ CB · ยังไม่ยืนยันผลราย Order" if template else (
          "สำเร็จตาม Check ใน Sheet" if checked else
          (value(cells[6]) if any(k in value(cells[6]) for k in
          ("สำเร็จ","เสร็จ","อยู่ระหว่าง","รอ","ยังไม่")) else "ยังไม่พบผลตรวจรับใน Sheet"))
        rows.append({"code":code,"row":rn,"sheet":sheet.title,
          "task":safe(cells[4],650),"owner":safe(cells[3],180),
          "due":safe(cells[5],100),"status":status,
          "check_raw":safe(cells[2],80),"notes":safe(cells[6],300),
          "process_group":safe(cells[0],220),"reference_only":template})
    return rows
def cb_tab_matches(sheet_name,record):
    t=re.sub(r"[^a-z0-9]","",sheet_name.casefold())
    order=re.sub(r"[^a-z0-9]","",record.get("reference","").casefold())
    no=re.sub(r"\D","",record["id"])
    if not order:return False
    d=re.sub(r"\D","",order)
    if len(d)>=4 and no and d[-4:] in t:
        return bool(re.search(r"(?:^|\D)"+re.escape(no)+r"(?:\D|$)",sheet_name))
    return order in t

# CB: all 88 real order lines, statuses are raw file values.
try:
    book=load_workbook(conf["sources"]["cb"]["path"],read_only=True,data_only=True,keep_links=False)
    tab=book["CB List"] if "CB List" in book.sheetnames else book.active
    records=[]
    template=book["ต้นฉบับ"] if "ต้นฉบับ" in book.sheetnames else None
    reference_steps=cb_sheet_steps(template,True) if template else []
    project_tabs=[sh for sh in book if sh.title not in ("CB List","ต้นฉบับ")]
    used_tabs=set()
    for row in tab.iter_rows(min_row=8,max_col=11,values_only=True):
        no=value(row[0])
        if not no.isdigit() or not row[1]:continue
        order=value(row[1]); st=value(row[10]);title=value(row[7]) if value(row[7])!="-" else value(row[8])
        rec=create("cb","CB-"+no,title,st,value(row[6]),value(row[5]),"วันที่โอน",
            value(row[9]),order,value(row[2]))
        rec["customer"]=safe(row[3],200)
        rec["product_category"]=safe(row[2],100)
        rec["cb_no"]=no
        match=[sh for sh in project_tabs if cb_tab_matches(sh.title,rec)]
        if match:
            steps=[]
            for sh in match:
                steps.extend(cb_sheet_steps(sh,False))
                used_tabs.add(sh.title)
            rec["steps"]=steps
            rec["cb_detail_origin"]="order_sheets"
            rec["cb_sheet_names"]=[sh.title for sh in match]
        elif reference_steps:
            category=rec["product_category"].casefold()
            selected="2.2" if any(x in category for x in ("ics","icf","edu")) else "2.1"
            rec["steps"]=[dict(x) for x in reference_steps
               if not x["code"].startswith("2.") or x["code"]==selected]
            rec["cb_detail_origin"]="template_reference_only"
            rec["cb_sheet_names"]=["ต้นฉบับ"]
        else:
            rec["steps"]=[]
            rec["cb_detail_origin"]="no_substage_sheet"
            rec["cb_sheet_names"]=[]
        records.append(rec)
    unused=[sh.title for sh in project_tabs if sh.title not in used_tabs]
    if unused:
        warnings.append({"source":"cb","error":"CB_SHEETS_UNLINKED","sheets":unused[:60]})
    book.close()
    groups["cb"]=records
except Exception as e:warnings.append({"source":"cb","error":type(e).__name__})
# ICS, ICF and ICE: gate rows. Column H is a recorded work date, not a confirmed deadline.
for typ in ("ics","icf","ice"):
    try:
        book=load_workbook(conf["sources"][typ]["path"],read_only=True,data_only=True,keep_links=False)
        tab=book["ICS69" if typ=="ics" else "ICF 2569"]
        records=[];parent="";ref="";typelabel=""
        for line,row in enumerate(tab.iter_rows(min_row=11,max_col=11,values_only=True),11):
            a=list(row)+[None]*max(0,11-len(row))
            if a[2]:parent=safe(a[2],240)
            if a[4]:ref=safe(a[4],140)
            if a[0]:typelabel=safe(a[0],100)
            gate=value(a[5]);owner=value(a[6]);st=value(a[8])
            if not gate or gate.lower()=="status":continue
            details=(value(a[9])+" "+value(a[10])).strip()
            rec=create(typ,typ.upper()+"-"+str(line),gate,st,owner,value(a[7]),"วันที่ในตาราง (ไม่ใช่ Due)",
                details,ref,parent)
            rec["work_type"]=typelabel
            records.append(rec)
        groups[typ]=records
        book.close()
    except Exception as e:warnings.append({"source":typ,"error":type(e).__name__})
# Innovation: file downloaded from authenticated SharePoint. Can be stale;
# label this separately from current cloud headline. Never present as live Graph data.
try:
    candidates=[x for x in (pathlib.Path.home()/"Downloads").glob("*.xlsx")
                if "นวัตกรรม" in x.name and x.stat().st_size>1000000]
    if candidates:
        path=max(candidates,key=lambda x:x.stat().st_mtime)
        book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
        records=[]
        for sheet_name,block in (("ติดตามงาน","ติดตามงาน"),("งานที่จบแล้ว","งานที่จบแล้ว")):
            if sheet_name not in book.sheetnames:continue
            sh=book[sheet_name];mgr="";cur=None;count=0
            def flush():
                if not cur:return
                steps=cur.pop("_steps")
                open_steps=[s for s in steps if s["status"]!="สำเร็จ"]
                explicit=sum(1 for s in open_steps if s["status"]=="กำลังดำเนินการ")
                cur["status"]="มีขั้นตอนค้าง "+str(len(open_steps)) if block=="ติดตามงาน" and open_steps else ("อยู่ในชีตงานที่จบแล้ว" if block!="ติดตามงาน" else "สำเร็จตามขั้นตอน")
                cur["steps"]=steps
                cur["notes"]="ตรวจขั้นตอนย่อย "+str(len(steps))+" ขั้น · ค้าง "+str(len(open_steps))+" ขั้น (บางช่องว่างยังต้องตรวจ)"
                cur["progress"]={"open_steps":len(open_steps),"active_steps":explicit,"total_steps":len(steps)}
                records.append(cur)
            for rn,row in enumerate(sh.iter_rows(min_row=1,max_col=10,values_only=True),1):
                v=list(row)+[None]*max(0,10-len(row))
                a=value(v[0])
                if "ผู้จัดการโครงการ" in a:mgr=a.split(":")[-1].strip()
                if a.isdigit() and (v[2] or v[1]):
                    flush()
                    cur=create("innovation",("INNO-" if block=="ติดตามงาน" else "CLOSED-")+str(rn),value(v[2]) or value(v[1]),"",mgr,"","ไม่ระบุ",sheet_name,
                        a,block)
                    cur["_steps"]=[]
                if cur and v[3] and value(v[3])!="แผนงาน":
                    cur["_steps"].append({"row":rn,"task":safe(v[3],350),"owner":safe(v[4],160),"due":safe(v[5],100),"status":safe(v[6],120) or "ไม่ระบุ","notes":safe(v[9],500)})
            flush()
        book.close()
        groups["innovation"]=records
        origin_time=dt.datetime.fromtimestamp(path.stat().st_mtime,dt.timezone.utc).isoformat()
    else:warnings.append({"source":"innovation","error":"No authenticated local workbook snapshot"})
except Exception as e:warnings.append({"source":"innovation","error":type(e).__name__})
# Planner: authenticated connector export on NB; assigned IDs are not automatically human names.
try:
    p=PRIVATE/"planner-export.json"
    exported=json.loads(p.read_text(encoding="utf-8-sig"))
    records=[]
    for t in exported.get("tasks",[]):
        stamp=t.get("due","")
        if stamp:
            try:due=dt.datetime.fromisoformat(stamp.replace("Z","+00:00")).astimezone(dt.timezone(dt.timedelta(hours=7))).date().isoformat()
            except ValueError:due=stamp
        else:due=""
        ass=t.get("owner_ids",[])
        owner=("Assignment "+str(len(ass))+" ราย (ต้องตรวจชื่อใน Planner)") if ass else "ไม่มี Assignment ใน Planner"
        records.append(create("planner",t["id"],t["title"],t["status"],owner,due,"กำหนดส่ง","ความคืบหน้า "+str(t["percent"])+"%",t["id"],"",t["percent"]))
    groups["planner"]=records
    planner_time=exported.get("retrieved_at")
except Exception as e:warnings.append({"source":"planner","error":type(e).__name__})
result={"schema":1,"generated_at":dt.datetime.now(dt.timezone.utc).isoformat(),
        "sources":{k:{"count":len(items),"rows":items} for k,items in groups.items()},
        "source_snapshots":{"innovation_downloaded":locals().get("origin_time"),"planner_exported":locals().get("planner_time")},
        "warnings":warnings,"security":"LOOPBACK_ONLY_DO_NOT_PUBLISH"}
dest=PRIVATE/"records.json"
tmp=PRIVATE/"records.json.tmp"
tmp.write_text(json.dumps(result,ensure_ascii=False,indent=1),encoding="utf8")
tmp.replace(dest)
print(json.dumps({"private_output":str(dest),"counts":{k:len(v) for k,v in groups.items()},"warnings":warnings},ensure_ascii=True))
